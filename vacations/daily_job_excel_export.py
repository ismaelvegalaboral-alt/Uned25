from collections import defaultdict
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.http import HttpResponse
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


TEMPLATE_PATH = (
    Path(settings.BASE_DIR)
    / "vacations"
    / "static"
    / "vacations"
    / "excel_templates"
    / "plantilla_faena.xlsx"
)

MONTHS_ES = [
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
]

WEEKDAYS_ES = [
    "LUNES", "MARTES", "MI?RCOLES", "JUEVES", "VIERNES", "S?BADO", "DOMINGO",
]


def _clean(value):
    return "" if value is None else str(value).strip()


def _anchor_cell(ws, row, col):
    cell = ws.cell(row=row, column=col)

    if not isinstance(cell, MergedCell):
        return cell

    coord = cell.coordinate
    for merged_range in ws.merged_cells.ranges:
        if coord in merged_range:
            return ws.cell(row=merged_range.min_row, column=merged_range.min_col)

    return cell


def _write(ws, row, col, value):
    cell = _anchor_cell(ws, row, col)
    if not isinstance(cell, MergedCell):
        cell.value = value


def _clear_template(ws):
    # Mantiene estilos, bordes, colores, alturas y anchuras.
    # Evita escribir en celdas combinadas que no son ancla.
    for row in range(1, 39):
        for col in range(1, 13):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell, MergedCell):
                continue
            cell.value = None


def _resource_text(item):
    parts = []

    if item.machine:
        parts.append(item.machine)

    if item.truck:
        parts.append(item.truck)

    if not parts and item.category and item.category != "work":
        parts.append(item.get_category_display())

    return " + ".join(parts)


def _worker_text(item):
    if item.worker_name:
        return item.worker_name
    if item.employee:
        return item.employee.full_name
    return ""


def _write_date(ws, plan_date):
    # Ajuste base para la cabecera de la plantilla.
    _write(ws, 1, 4, WEEKDAYS_ES[plan_date.weekday()])   # D1
    _write(ws, 1, 5, plan_date.day)                      # E1
    _write(ws, 1, 6, MONTHS_ES[plan_date.month - 1])     # F1
    _write(ws, 1, 8, plan_date.year)                     # H1


def _group_assignments(assignments):
    groups = defaultdict(list)

    for item in assignments:
        key = (
            _clean(item.client) or "SIN CLIENTE",
            _clean(item.worksite) or "SIN OBRA/ZONA",
        )
        groups[key].append(item)

    return list(groups.items())


def _write_work_blocks(ws, assignments):
    # Tres bloques visuales como en el Excel original: A:B, D:E y G:H.
    blocks = [
        {"cols": (1, 2), "start": 3, "end": 38},
        {"cols": (4, 5), "start": 3, "end": 38},
        {"cols": (7, 8), "start": 3, "end": 38},
    ]

    groups = _group_assignments(assignments)

    block_index = 0
    row = blocks[block_index]["start"]

    for (client, worksite), items in groups:
        if block_index >= len(blocks):
            break

        block = blocks[block_index]
        col_a, col_b = block["cols"]

        needed = 1 + max(1, len(items)) + 1
        if row + needed - 1 > block["end"]:
            block_index += 1
            if block_index >= len(blocks):
                break

            block = blocks[block_index]
            col_a, col_b = block["cols"]
            row = block["start"]

        _write(ws, row, col_a, client)
        _write(ws, row, col_b, worksite)
        row += 1

        for item in items:
            if row > block["end"]:
                block_index += 1
                if block_index >= len(blocks):
                    return

                block = blocks[block_index]
                col_a, col_b = block["cols"]
                row = block["start"]

            resource = _resource_text(item)
            worker = _worker_text(item)

            extras = []
            if item.hours:
                extras.append(f"{item.hours} h")
            if item.trips:
                extras.append(f"{item.trips} viajes")
            if item.notes:
                extras.append(item.notes)

            if extras:
                resource = f"{resource} ({' ? '.join(extras)})" if resource else " ? ".join(extras)

            _write(ws, row, col_a, resource)
            _write(ws, row, col_b, worker)
            row += 1

        row += 1


def _write_status_blocks(ws, status_entries):
    # Zona derecha K:L de la plantilla.
    status_by_type = defaultdict(list)

    for entry in status_entries:
        status_by_type[entry.get_status_type_display()].append(entry.text)

    cols = [11, 12]  # K, L
    col_index = 0
    row = 3

    for status_label, texts in status_by_type.items():
        col = cols[col_index]

        if row > 37:
            if col_index == 0:
                col_index = 1
                row = 3
                col = cols[col_index]
            else:
                break

        _write(ws, row, col, status_label.upper())
        row += 1

        for text in texts:
            if row > 38:
                if col_index == 0:
                    col_index = 1
                    row = 3
                    col = cols[col_index]
                else:
                    return

            _write(ws, row, col, text)
            row += 1

        row += 1


def build_daily_job_plan_excel_response(plan):
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active

    _clear_template(ws)
    _write_date(ws, plan.plan_date)

    assignments = list(plan.assignments.select_related("employee").all())
    status_entries = list(plan.status_entries.all())

    _write_work_blocks(ws, assignments)
    _write_status_blocks(ws, status_entries)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"faena_{plan.plan_date:%Y%m%d}.xlsx"

    response = HttpResponse(
        output.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
